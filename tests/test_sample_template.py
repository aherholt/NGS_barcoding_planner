"""Excel sample-list template: download, fill in, upload (round trip)."""

import io

from django.urls import reverse
from openpyxl import load_workbook

from planner import services
from planner.sample_template import build
from planner.tabular import read_table

from .conftest import sample_rows


def test_generic_template_structure(db):
    wb = load_workbook(io.BytesIO(build()))
    assert wb.sheetnames == ["Samples", "Instructions", "Lists"]
    ws = wb["Samples"]
    headers = [c.value for c in ws[1]]
    assert headers == ["sample_id", "source_plate", "source_well", "plate_format", "condition", "pool_id", "barcode_id"]
    validated = {str(dv.sqref).split(":")[0] for dv in ws.data_validations.dataValidation}
    assert {"A2", "C2", "D2"} <= validated  # unique ids, wells, plate format
    assert ws["A2"].number_format == "@"
    assert read_table(build(), "t.xlsx") == []  # empty template uploads nothing


def test_fill_in_and_upload(tp_experiment, users):
    wb = load_workbook(io.BytesIO(build(tp_experiment)))
    wb.move_sheet("Instructions", offset=-1)  # user reorders sheets → upload still finds "Samples"
    ws = wb["Samples"]
    for r, (sid, well) in enumerate((("001", "A1"), ("002", "B1"), ("003", "A2")), start=2):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = sid, "CULT01", well, 6, "DMSO"
    buf = io.BytesIO()
    wb.save(buf)
    rows = read_table(buf.getvalue(), "filled.xlsx")
    assert services.upload_samples(tp_experiment, rows, users["tech"]) == 3
    s = tp_experiment.samples.get(sample_id="001")
    assert (s.source_well, s.plate_format) == ("A1", 6)


def test_round_trip_keeps_plan(tp_experiment, users):
    services.upload_samples(tp_experiment, sample_rows(30), users["tech"])
    services.plan_barcodes(tp_experiment, seed=2)
    before = {s.sample_id: (s.pool.pool_id, s.well_barcode.barcode_id) for s in tp_experiment.samples.all()}
    rows = read_table(build(tp_experiment), "x.xlsx")  # pre-filled with current samples incl. plan
    assert len(rows) == 30
    services.upload_samples(tp_experiment, rows, users["tech"])
    after = {s.sample_id: (s.pool.pool_id, s.well_barcode.barcode_id) for s in tp_experiment.samples.all()}
    assert after == before and services.validate_plan(tp_experiment) == []
    lists = load_workbook(io.BytesIO(build(tp_experiment)))["Lists"]
    assert lists["C2"].value == "BC01"  # barcode drop-down from the experiment's set


def test_bulk_template_has_no_plan_columns(db):
    from planner.models import Experiment
    e = Experiment.objects.create(code="RNA-T", title="bulk", library_type="bulk_rnaseq")
    headers = [c.value for c in load_workbook(io.BytesIO(build(e)))["Samples"][1]]
    assert "pool_id" not in headers and headers[0] == "sample_id"


def test_download_views(client, users, tp_experiment):
    client.force_login(users["tech"])
    r = client.get(reverse("experiment_sample_template", args=[tp_experiment.pk]))
    assert r.status_code == 200 and "TP-1_sample_list.xlsx" in r["Content-Disposition"]
    assert client.get(reverse("sample_template")).status_code == 200
    page = client.get(reverse("experiment_detail", args=[tp_experiment.pk])).content.decode()
    assert "Excel template" in page
