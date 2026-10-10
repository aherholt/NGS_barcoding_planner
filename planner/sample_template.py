"""Excel template for the sample list (download → fill in → upload).

Sheet "Samples" is what the upload reads (column names must stay as they are).
Sheet "Instructions" explains every column; sheet "Lists" holds the drop-down values.
If the experiment already has samples, they are pre-filled, so the template doubles as
"download current list, edit, upload again".
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import plates

MAX_ROWS = 2000
TEXT_ROWS = 500  # rows pre-formatted as text
REQUIRED_FILL = PatternFill("solid", fgColor="F4A6A6")
OPTIONAL_FILL = PatternFill("solid", fgColor="D9E2F3")
PLAN_FILL = PatternFill("solid", fgColor="EDEDED")

# (column, required?, width, description, allowed values, example)
COLUMNS = [
    ("sample_id", True, 18, "Unique name of the sample (no duplicates).", "free text", "P1_A1"),
    ("source_plate", None, 14, "ID of the culture plate (or tube rack) the lysate comes from.", "free text", "CULT01"),
    ("source_well", None, 12, "Well in that plate, column-wise A1, B1, C1 …", "must exist in the plate format", "A1"),
    ("plate_format", False, 13, "Culture plate format. Empty = 24. All rows of one plate must have the same format.",
     "6, 12, 24, 48, 96", "24"),
    ("condition", False, 28, "Treatment / genotype / replicate — kept in the barcode map for analysis.", "free text", "DMSO_rep1"),
]
PLAN_COLUMNS = [
    ("pool_id", False, 16, "Optional: import a ready-made plan. Fill pool_id AND barcode_id in every row, "
     "otherwise leave both empty and use 'Plan pools & barcodes' in the app.", "free text", "TP26-001_P01"),
    ("barcode_id", False, 12, "Optional, with pool_id: well barcode from the experiment's barcode set; "
     "each barcode only once per pool.", "barcode IDs of the set", "BC01"),
]


def build(experiment=None) -> bytes:
    tag_and_pool = experiment is None or experiment.is_tag_and_pool
    robot = tag_and_pool  # Tag&Pool: plate/well are needed by the STAR robot
    cols = COLUMNS + (PLAN_COLUMNS if tag_and_pool else [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Samples"
    lists = wb.create_sheet("Lists")
    info = wb.create_sheet("Instructions")
    wb.move_sheet(info, offset=-1)  # order: Samples, Instructions, Lists
    lists.sheet_state = "hidden"

    # --- drop-down source lists -------------------------------------------------
    lists["A1"] = "plate_format"
    for i, f in enumerate(plates.FORMATS, start=2):
        lists.cell(row=i, column=1, value=f)
    lists["B1"] = "wells (96-well, covers all formats)"
    for i, w in enumerate(plates.wells(96), start=2):
        lists.cell(row=i, column=2, value=w)
    barcodes = []
    if tag_and_pool and experiment is not None and experiment.well_barcode_set:
        barcodes = list(experiment.well_barcode_set.barcodes.values_list("barcode_id", flat=True))
    lists["C1"] = "barcode_id"
    for i, b in enumerate(barcodes, start=2):
        lists.cell(row=i, column=3, value=b)

    # --- header -------------------------------------------------------------------
    for c, (name, required, width, desc, allowed, _ex) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        req = required or (required is None and robot)
        cell.fill = REQUIRED_FILL if req else (PLAN_FILL if name in ("pool_id", "barcode_id") else OPTIONAL_FILL)
        cell.font = Font(bold=True)
        cell.comment = Comment(f"{'REQUIRED. ' if req else 'optional. '}{desc}\nAllowed: {allowed}", "NGS planner")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"
    # text columns stay text (Excel would turn "001" into 1 or "1-2" into a date)
    for c, (name, *_rest) in enumerate(cols, start=1):
        if name != "plate_format":
            for r in range(2, TEXT_ROWS + 2):
                ws.cell(row=r, column=c).number_format = "@"
    col_of = {name: get_column_letter(i) for i, (name, *_rest) in enumerate(cols, start=1)}

    # --- drop-downs (invalid entries are also rejected by the app on upload) -------
    def validation(formula, column, title, text):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorTitle=title, error=text)
        dv.add(f"{col_of[column]}2:{col_of[column]}{MAX_ROWS}")
        ws.add_data_validation(dv)

    validation(f"=Lists!$A$2:$A${len(plates.FORMATS) + 1}", "plate_format", "Plate format", "Choose 6, 12, 24, 48 or 96.")
    validation("=Lists!$B$2:$B$97", "source_well", "Well", "Use a well name like A1 … H12.")
    if barcodes:
        validation(f"=Lists!$C$2:$C${len(barcodes) + 1}", "barcode_id", "Barcode",
                   f"Choose a barcode of {experiment.well_barcode_set}.")
    dup = DataValidation(type="custom", formula1=f"COUNTIF($A$2:$A${MAX_ROWS},A2)=1", allow_blank=True,
                         showErrorMessage=True, errorTitle="Duplicate", error="sample_id must be unique.")
    dup.add(f"A2:A{MAX_ROWS}")
    ws.add_data_validation(dup)

    # --- pre-fill existing samples ----------------------------------------------------
    if experiment is not None:
        for r, s in enumerate(experiment.samples.select_related("pool", "well_barcode"), start=2):
            values = {"sample_id": s.sample_id, "source_plate": s.source_plate, "source_well": s.source_well,
                      "plate_format": s.plate_format, "condition": s.condition,
                      "pool_id": s.pool.pool_id if (s.pool and tag_and_pool) else None,
                      "barcode_id": s.well_barcode.barcode_id if s.well_barcode else None}
            for name, letter in col_of.items():
                ws[f"{letter}{r}"] = values.get(name)

    # --- instructions -------------------------------------------------------------------
    title = f"Sample list for {experiment.code} – {experiment.title}" if experiment else "Sample list template"
    info["A1"] = title
    info["A1"].font = Font(bold=True, size=14)
    lines = [
        "1. Fill in the sheet 'Samples': one row per sample (lysate). Do not rename the columns or the sheet.",
        "2. Red columns are required" + (" (the STAR robot needs culture plate and well)." if robot else "."),
        "   Blue columns are optional; grey columns only if you import a ready-made pool/barcode plan.",
        "3. Save as .xlsx (or CSV) and upload it on the experiment page: 'Upload sample list'.",
        "   The upload replaces the current sample list; the app checks everything and lists all problems at once.",
        "4. Wells are counted column-wise (A1, B1, C1, …, A2 …), the order in which the robot works through a plate.",
    ]
    if experiment is not None and experiment.samples.exists():
        lines.append("Note: this file was pre-filled with the current sample list of the experiment.")
    for i, t in enumerate(lines, start=3):
        info[f"A{i}"] = t
    start = len(lines) + 5
    for c, h in enumerate(("Column", "Required", "Meaning", "Allowed values", "Example"), start=1):
        cell = info.cell(row=start, column=c, value=h)
        cell.font = Font(bold=True)
    for r, (name, required, _w, desc, allowed, example) in enumerate(cols, start=start + 1):
        req = required or (required is None and robot)
        info.cell(row=r, column=1, value=name).font = Font(bold=True)
        info.cell(row=r, column=2, value="yes" if req else "no")
        info.cell(row=r, column=3, value=desc).alignment = Alignment(wrap_text=True, vertical="top")
        info.cell(row=r, column=4, value=allowed)
        info.cell(row=r, column=5, value=example)
        for c in range(1, 6):
            info.cell(row=r, column=c).alignment = Alignment(wrap_text=(c == 3), vertical="top")
    r = start + len(cols) + 2
    info.cell(row=r, column=1, value="Plate formats (rows × columns):").font = Font(bold=True)
    for i, (f, (rows, columns)) in enumerate(plates.FORMATS.items(), start=r + 1):
        info.cell(row=i, column=1, value=f"{f}-well")
        info.cell(row=i, column=2, value=f"{rows} × {columns}")
        info.cell(row=i, column=3, value=f"A1 … {plates.wells(f)[-1]}")
    if experiment is not None and experiment.well_barcode_set:
        info.cell(row=r + 7, column=1, value=f"Well-barcode set: {experiment.well_barcode_set} "
                                             f"({experiment.well_barcode_set.size} barcodes = max. samples per pool)")
    for col, w in zip("ABCDE", (16, 10, 70, 26, 16)):
        info.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
