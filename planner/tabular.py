"""Read CSV or Excel (.xlsx) files into a list of dicts with lower-case column names."""

from __future__ import annotations

import io
from pathlib import Path

from .services import read_csv_text


def read_table(data: bytes, filename: str, sheet: str | None = None) -> list[dict]:
    if Path(filename).suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h or "").strip().lower() for h in rows[0]]
        out = []
        for r in rows[1:]:
            if all(v in (None, "") for v in r):
                continue
            out.append({h: ("" if v is None else str(v).strip()) for h, v in zip(header, r) if h})
        return out
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return read_csv_text(data.decode(enc))
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not read the file (unknown text encoding).")
